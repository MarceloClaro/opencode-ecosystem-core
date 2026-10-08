"""Regressões de cálculos e preenchimento. Executar com Python e python-docx."""
import copy
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from docx import Document
from gerar_documentos import (analyze, derive, fill, money, replace_spanning,
                              validate_input, xml_text)


class GeradorTests(unittest.TestCase):
    def base(self, fields=None, lists=None):
        fields = fields or {}
        return {'campos':fields,'fontes':{k:'Fonte sintética de teste' for k in fields},
                'listas':lists or {}}

    def test_money_exact_and_ambiguous(self):
        for value in ['1.250,50','1250,50','1250.50']:
            self.assertEqual(money(value),Decimal('1250.50'))
        for value in ['1.250','NaN','Infinity','1,234,56',1.2]:
            with self.subTest(value=value), self.assertRaises(ValueError):money(value)

    def test_rounding_and_immutability(self):
        raw=self.base(lists={'bens':[{'quantidade':'0.5','unitario':'0.01','_fonte':'Teste'}]})
        snapshot=copy.deepcopy(raw)
        data,_=derive(raw)
        self.assertEqual(data['campos']['bens.total'],'0,01')
        self.assertEqual(raw,snapshot)

    def test_partial_and_empty_are_not_totals(self):
        for rows in [[],[{'quantidade':'2','unitario':None,'_fonte':'Teste'}],
                     [{'quantidade':'2','unitario':'10'}]]:
            data,_=derive(self.base({'bens.total':'20'}, {'bens':rows}))
            self.assertIsNone(data['campos']['bens.total'])

    def test_product_mismatch_blocks(self):
        with self.assertRaises(ValueError):
            derive(self.base(lists={'bens':[{'quantidade':'2','unitario':'10','total':'19','_fonte':'Teste'}]}))

    def test_action_balance_and_conflict(self):
        row={'programado':'50','executado':'12.50','_fonte':'Teste'}
        data,_=derive(self.base(lists={'acoes':[row]}))
        self.assertEqual(data['campos']['acoes.saldo'],'37,50')
        row['saldo']='12.50'
        with self.assertRaises(ValueError):derive(self.base(lists={'acoes':[row]}))

    def test_unknown_opening_balance_does_not_derive(self):
        data,_=derive(self.base({'financeiro.repasse':'100','financeiro.contrapartida':'0',
                                'financeiro.rendimentos':'0','financeiro.despesas':'10'}))
        self.assertNotIn('financeiro.saldo',data['campos'])

    def test_zero_is_known_and_financial_balance_exact(self):
        data,_=derive(self.base({'financeiro.saldo_anterior':'0','financeiro.repasse':'100',
                                'financeiro.contrapartida':'0','financeiro.rendimentos':'0.10',
                                'financeiro.despesas':'90.10'}))
        self.assertEqual(data['campos']['financeiro.saldo'],'10,00')

    def test_payments_cross_check(self):
        with self.assertRaises(ValueError):
            derive(self.base({'financeiro.despesas':'10'},
                {'pagamentos':[{'valor':'9','_fonte':'Teste'}]}))

    def test_duplicate_alert_keeps_rows(self):
        row={'comprovante':'TESTE-1','credor':'Teste','data_pagamento':'01/01/2000',
             'valor':'10','_fonte':'Teste'}
        data=self.base(lists={'pagamentos':[dict(row),dict(row)]})
        self.assertEqual(analyze(data)[0]['codigo'],'possivel_duplicidade')
        self.assertEqual(len(data['listas']['pagamentos']),2)

    def test_selected_quote_alert(self):
        data=self.base(lists={'mapa':[{'escolhido':'A','a':'10','unitario':'11','_fonte':'Teste'}]})
        self.assertEqual(analyze(data)[0]['codigo'],'preco_escolhido_diverge')

    def test_preflight_rejects_invalid_prices(self):
        with self.assertRaises(ValueError):
            validate_input(self.base(lists={'mapa':[{'a':'valor inválido'}]}))

    def test_marker_across_runs(self):
        doc=Document();p=doc.add_paragraph()
        p.add_run('Antes {{escola.');p.add_run('nome}} depois')
        replace_spanning(p._p,{'escola.nome':'Exemplo'})
        self.assertEqual(p.text,'Antes Exemplo depois')

    def test_expansion_missing_sources_and_currency(self):
        rows=[{'numero':str(i),'nota':f'TESTE-{i}','quantidade':'2','unitario':'12.50',
               'descricao':'Descrição sintética','_fonte':'Teste'} for i in range(1,6)]
        data,_=derive(self.base({'uex.nome':'Nome sem fonte'},{'bens':rows}))
        data['fontes'].pop('uex.nome')
        with tempfile.TemporaryDirectory() as folder:
            report=fill('anexo-05',data,Path(folder))
            text=xml_text(Document(Path(folder)/report['arquivo'])._element)
        self.assertNotIn('{{',text)
        self.assertNotIn('Nome sem fonte',text)
        self.assertIn('[PENDENTE]',text)
        self.assertIn('12,50',text)
        self.assertIn('125,00',text)
        for i in range(1,6):self.assertIn(f'TESTE-{i}',text)


if __name__ == '__main__':
    unittest.main()
