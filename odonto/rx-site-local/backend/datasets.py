"""Inventário e amostra REAIS dos datasets locais — SPEC-935-R900 FR-003/FR-004.

Somente leitura. Nenhum download, nenhuma escrita. ``sha256`` do Table_S1.xlsx
é conferido contra o valor fixado no manuscrito OdontoCA
(``odontoca.data.SOURCE_SHA256``); divergência é reportada, nunca silenciada.
"""
from __future__ import annotations

import hashlib
import math
from pathlib import Path
from typing import Any

from .paths import DATASETS, TABLE_S1, ensure_import_path, resolve_dataset

ensure_import_path()

SOURCE_SHA256 = "b7819fee81efbe2e227b7700c3cd86ce8bc6f7fe6f49da6214f4107d099184fa"

_IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}
_TEXT_SUFFIXES = {".csv", ".tsv", ".txt", ".json"}
_CACHE: dict[str, Any] | None = None


def sha256_file(path: Path, chunk: int = 1 << 20) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(chunk), b""):
            digest.update(block)
    return digest.hexdigest()


def _to_jsonable(value: Any) -> Any:
    """Converte tipos numpy/pandas para JSON puro; NaN/inf → None."""
    if value is None:
        return None
    if isinstance(value, float):
        return None if (math.isnan(value) or math.isinf(value)) else value
    if isinstance(value, (str, bool, int)):
        return value
    if hasattr(value, "item"):
        try:
            return _to_jsonable(value.item())
        except Exception:
            return str(value)
    if hasattr(value, "isoformat"):
        return value.isoformat()
    try:
        if isinstance(value, float) and math.isnan(value):
            return None
    except Exception:
        pass
    return str(value)


def _count_rows_text(path: Path) -> int:
    with path.open("rb") as fh:
        total = sum(1 for _ in fh)
    return max(0, total - 1)


def _text_header(path: Path) -> list[str]:
    with path.open("r", encoding="utf-8", errors="replace") as fh:
        first = fh.readline().strip("\n\r")
    sep = "\t" if path.suffix.lower() == ".tsv" else ","
    return [c.strip().strip('"') for c in first.split(sep)][:80]


def _xlsx_sheets(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        return [{"name": ws.title, "rows": int(ws.max_row or 0)} for ws in wb.worksheets]
    except Exception:
        return []


def describe_file(path: Path, name: str) -> dict[str, Any]:
    suffix = path.suffix.lower()
    entry: dict[str, Any] = {
        "name": name,
        "path": str(path.relative_to(DATASETS.parent)) if path.is_relative_to(DATASETS) else str(path),
        "kind": "xlsx" if suffix in {".xlsx", ".xls"} else (
            "image" if suffix in _IMAGE_SUFFIXES else "text" if suffix in _TEXT_SUFFIXES else "file"
        ),
        "bytes": path.stat().st_size,
        "sha256": sha256_file(path),
        "source": "local",
    }
    if suffix in {".csv", ".tsv", ".txt"}:
        entry["rows"] = _count_rows_text(path)
        entry["columns"] = _text_header(path)
    elif suffix in {".xlsx", ".xls"}:
        entry["sheets"] = _xlsx_sheets(path)
        entry["sha256_expected"] = SOURCE_SHA256
        entry["sha256_ok"] = entry["sha256"] == SOURCE_SHA256
    return entry


def describe_dir(path: Path, name: str) -> dict[str, Any]:
    files = sorted(p for p in path.rglob("*") if p.is_file())
    total = sum(p.stat().st_size for p in files)
    images = [p for p in files if p.suffix.lower() in _IMAGE_SUFFIXES]
    texts = [p for p in files if p.suffix.lower() in {".csv", ".tsv"}]
    children: list[dict[str, Any]] = []
    for p in texts:
        try:
            children.append(describe_file(p, f"{name}/{p.relative_to(path)}"))
        except Exception:
            continue
    return {
        "name": name,
        "path": str(path.relative_to(DATASETS.parent)),
        "kind": "dir",
        "bytes": int(total),
        "files": len(files),
        "images": len(images),
        "texts": len(texts),
        "sha256": None,
        "source": "local",
        "children": children,
        "sample_files": [str(p.relative_to(path)) for p in files[:5]],
    }


def inventory(refresh: bool = False) -> list[dict[str, Any]]:
    """Lista real de datasets + Table_S1.xlsx (fonte primária). Ordenada."""
    global _CACHE
    if _CACHE is not None and not refresh:
        return _CACHE["datasets"]

    datasets: list[dict[str, Any]] = []
    if DATASETS.is_dir():
        for path in sorted(DATASETS.iterdir(), key=lambda p: p.name.lower()):
            if path.name.startswith("."):
                continue
            try:
                if path.is_dir():
                    datasets.append(describe_dir(path, path.name))
                else:
                    datasets.append(describe_file(path, path.name))
            except Exception as exc:  # nunca derruba o inventário inteiro
                datasets.append({"name": path.name, "kind": "error", "error": str(exc)})

    if TABLE_S1.is_file():
        entry = describe_file(TABLE_S1, "Table_S1.xlsx")
        entry["path"] = str(TABLE_S1.relative_to(TABLE_S1.parents[2]))
        entry["role"] = "fonte_primaria_coorte"
        datasets.append(entry)

    _CACHE = {"datasets": datasets}
    return datasets


def names() -> set[str]:
    """Allowlist de nomes aceitos por ``/api/datasets/sample`` e ``/api/train``."""
    allowed: set[str] = set()
    for entry in inventory():
        if entry.get("kind") == "error":
            continue
        allowed.add(entry["name"])
        for child in entry.get("children", []):
            allowed.add(child["name"])
    return allowed


def sample(name: str, n: int = 5, sheet: str | None = None) -> dict[str, Any]:
    """Amostra real de um dataset do inventário. Fail-closed fora da allowlist."""
    if name not in names():
        raise KeyError(name)
    n = max(1, min(int(n), 50))
    if name == "Table_S1.xlsx":
        path = TABLE_S1
    else:
        path = resolve_dataset(name)

    import pandas as pd

    if path.suffix.lower() in {".xlsx", ".xls"}:
        frame = pd.read_excel(path, sheet_name=sheet or 0, nrows=n)
    else:
        try:
            frame = pd.read_csv(path, nrows=n)
        except UnicodeDecodeError:
            frame = pd.read_csv(path, nrows=n, encoding="latin-1")
    columns = [str(c) for c in frame.columns]
    rows = [{str(k): _to_jsonable(v) for k, v in record.items()} for record in frame.to_dict("records")]
    return {"name": name, "n": len(rows), "columns": columns, "rows": rows, "real": True}
