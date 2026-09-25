#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gerador DOCX v2.0 — PROTOCOLO (rascunho com pendências críticas)
SPEC-935-R522 — Educação Por Escrito Chamada 799
Correção anti-overclaim: remove todos os números/figuras de extração simulada;
preserva arquitetura científica e converte achados não rastreáveis em campos pendentes.
ABNT: A4, margens 3/2 cm, Times 12, 1.5, capa, resumo, abstract, seções.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUTPUT = ("/home/marceloclaro/opencode-ecosystem-core/manuscrito_porescrito799_R522/"
          "manuscrito_porescrito799_brasil_comparado_R522_v20_protocolo_ABNT.docx")

def set_margins(section, top=3, bottom=2, left=3, right=2):
    section.top_margin = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.left_margin = Cm(left)
    section.right_margin = Cm(right)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)

def set_paragraph_format(p, space_after=6, space_before=0, line_spacing=1.5,
                         alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_indent=Cm(1.25)):
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line_spacing
    pf.alignment = alignment
    if first_line_indent is not None:
        pf.first_line_indent = first_line_indent
    return p

def add_heading_custom(doc, text, level=1, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.bold = True
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    h.alignment = alignment
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.line_spacing = 1.0
    return h

def add_body(doc, text, bold=False, italic=False, space_after=6, first_line=True,
             alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, line_spacing=1.5):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    set_paragraph_format(p, space_after=space_after,
                         first_line_indent=Cm(1.25) if first_line else Cm(0),
                         alignment=alignment, line_spacing=line_spacing)
    return p

def add_center(doc, text, bold=False, italic=False, size=12, space_after=6, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    set_paragraph_format(p, space_after=space_after, first_line_indent=Cm(0),
                         alignment=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0)
    return p

def add_quote_long(doc, text):
    """Citação longa ABNT: recuo 4 cm, fonte 10, espaçamento simples."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(4)
    pf.space_after = Pt(6)
    pf.space_before = Pt(6)
    pf.line_spacing = 1.0
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    return p

def add_reference(doc, text):
    """Referência ABNT: sem recuo de 1ª linha, espaçamento simples."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.0
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    return p

def add_footer_with_page_number(doc):
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    p._p.append(fldChar1)
    p._p.append(instrText)
    p._p.append(fldChar2)
    run = p.add_run("  |  Educação Por Escrito — Chamada 799 — R522 — PROTOCOLO")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)

# ============================ DOCUMENTO ============================
doc = Document()
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
font.color.rgb = RGBColor(0, 0, 0)
pf = style.paragraph_format
pf.space_after = Pt(6)
pf.line_spacing = 1.5
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for section in doc.sections:
    set_margins(section, top=3, bottom=2, left=3, right=2)
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)

add_footer_with_page_number(doc)

# ============ CAPA ============
for _ in range(2):
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

add_center(doc, "STATUS: RASCUNHO ACADÊMICO COM PENDÊNCIAS CRÍTICAS",
           bold=True, size=13, space_after=2, color=RGBColor(0x8B, 0x1A, 0x1A))
add_center(doc, "NÃO SUBMETER NESTA VERSÃO", bold=True, size=12, space_after=14,
           color=RGBColor(0x8B, 0x1A, 0x1A))
add_body(doc,
         "Motivo: os arquivos-fonte identificam a triagem/extração Rayyan como simulada. "
         "Por integridade científica, os números de seleção, frequências, gap map e Figuras 2–3 "
         "não são tratados como resultados observados. Esta versão preserva a arquitetura científica, "
         "converte achados não rastreáveis em campos pendentes e separa proposições analíticas de "
         "evidência empírica.",
         size=9, italic=True, first_line=False, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)

add_center(doc, "PONTIFÍCIA UNIVERSIDADE CATÓLICA DO RIO GRANDE DO SUL\n"
                "PROGRAMA DE PÓS-GRADUAÇÃO EM EDUCAÇÃO", bold=True, size=12, space_after=12)

add_center(doc, "Educação Por Escrito — e-ISSN 2179-8435 — Qualis A4 (Ensino, 2021-2024)",
           italic=True, size=10, space_after=8, color=RGBColor(60, 60, 60))

add_center(doc, "CHAMADA 799 — FLUXO CONTÍNUO", bold=True, size=10, space_after=2)
add_center(doc, "https://revistaseletronicas.pucrs.br/porescrito/announcement/view/799",
           italic=True, size=8, space_after=18)

add_center(doc, "IA GENERATIVA NA EDUCAÇÃO JURÍDICA BRASILEIRA\n"
                "ENTRE A REGULAÇÃO E A SALA DE AULA", bold=True, size=14, space_after=4)
add_center(doc, "Protocolo e arquitetura de revisão de escopo Brasil-comparado sobre ética, "
                "governança algorítmica e equidade (2020–2025)", size=12, space_after=18)

add_center(doc, "Versão anônima para avaliação editorial futura — autoria e folha de "
                "identificação devem ser enviadas separadamente conforme a revista-alvo.",
           italic=True, size=10, space_after=20)

add_center(doc, "Artigo original — Revisão de Escopo (Scoping Review) — Protocolo PRISMA-ScR / JBI",
           italic=True, size=9, space_after=6)
add_center(doc, "Área: Ensino — Linha: Educação Jurídica — Meta de rigor: Qualis A1 "
                "(tratada como meta, não garantia)", size=8, space_after=12)

add_center(doc, "Porto Alegre — RS, 17 de setembro de 2026", size=10, space_after=12)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
pf = p.paragraph_format
pf.space_after = Pt(2)
run = p.add_run("Manuscrito v2.0 (PROTOCOLO) — formato ABNT — Times New Roman 12, 1,5, margens 3/2 cm — SPEC-935-R522\n")
run.font.name = 'Times New Roman'
run.font.size = Pt(7)
run.font.color.rgb = RGBColor(80, 80, 80)
run = p.add_run("Submissão: https://revistaseletronicas.pucrs.br/porescrito/about/submissions  |  Sem APC  |  Duplo-cega  |  Turnitin")
run.font.name = 'Times New Roman'
run.font.size = Pt(7)
run.font.color.rgb = RGBColor(80, 80, 80)
run.italic = True

doc.add_page_break()

# ============ RESUMO ============
add_heading_custom(doc, "RESUMO", level=1, alignment=WD_ALIGN_PARAGRAPH.CENTER)
add_body(doc,
         "Introdução. A difusão da inteligência artificial generativa (IAGen) na educação jurídica "
         "amplia possibilidades de apoio à aprendizagem, mas também intensifica questões de integridade "
         "acadêmica, proteção de dados, transparência, supervisão humana e equidade. Objetivo. Mapear a "
         "literatura científica publicada entre 2020 e 2025 sobre implementação ética, transparente e "
         "equitativa de IAGen na educação jurídica brasileira, em perspectiva comparada com União Europeia, "
         "Portugal, Alemanha e Estados Unidos. Método. Propõe-se revisão de escopo orientada pelo Joanna "
         "Briggs Institute e relatada segundo PRISMA-ScR, com buscas em bases multidisciplinares e "
         "educacionais, triagem independente, extração padronizada e síntese descritiva e temática. "
         "Resultados. [A CONFIRMAR APÓS TRIAGEM REAL E AUDITÁVEL]. Os números e gráficos presentes em "
         "versões anteriores não são reportados aqui porque os arquivos de origem os identificam como "
         "simulados. Conclusão. O protocolo permite testar, sem antecipação de achados, se a produção sobre "
         "educação jurídica e IAGen privilegia marcos normativos em detrimento de evidência pedagógica, "
         "governança e equidade.",
         size=10, first_line=False, space_after=8, line_spacing=1.0)

p = doc.add_paragraph()
p.paragraph_format.first_line_indent = Cm(0)
p.paragraph_format.space_after = Pt(12)
run = p.add_run("Palavras-chave: ")
run.bold = True
run.font.name = 'Times New Roman'
run.font.size = Pt(10)
run = p.add_run("inteligência artificial generativa; educação jurídica; ética; governança algorítmica; "
                "equidade; revisão de escopo.")
run.font.name = 'Times New Roman'
run.font.size = Pt(10)

# ============ ABSTRACT ============
add_heading_custom(doc, "ABSTRACT", level=1, alignment=WD_ALIGN_PARAGRAPH.CENTER)
add_body(doc,
         "Introduction. The diffusion of generative artificial intelligence (GenAI) in legal education "
         "creates opportunities for learning support while raising concerns about academic integrity, "
         "data protection, transparency, human oversight and equity. Objective. To map scientific "
         "literature published from 2020 to 2025 on the ethical, transparent and equitable implementation "
         "of GenAI in Brazilian legal education, with comparisons involving the European Union, Portugal, "
         "Germany and the United States. Method. A scoping review guided by the Joanna Briggs Institute and "
         "reported according to PRISMA-ScR is proposed, including multidisciplinary and educational "
         "databases, independent screening, standardized data extraction, and descriptive and thematic "
         "synthesis. Results. [TO BE CONFIRMED AFTER A REAL AND AUDITABLE SCREENING PROCESS]. Numbers and "
         "graphics from previous drafts are not reported because their source files identify them as "
         "simulated. Conclusion. The protocol is designed to test, without anticipating findings, whether "
         "scholarship on legal education and GenAI prioritizes regulatory frameworks over pedagogical "
         "evidence, governance and equity.",
         size=10, first_line=False, space_after=8, line_spacing=1.0)

p = doc.add_paragraph()
p.paragraph_format.first_line_indent = Cm(0)
p.paragraph_format.space_after = Pt(12)
run = p.add_run("Keywords: ")
run.bold = True
run.font.name = 'Times New Roman'
run.font.size = Pt(10)
run = p.add_run("generative artificial intelligence; legal education; ethics; algorithmic governance; "
                "equity; scoping review.")
run.font.name = 'Times New Roman'
run.font.size = Pt(10)

# ============ 1 INTRODUÇÃO ============
add_heading_custom(doc, "1 INTRODUÇÃO")
add_body(doc,
         "A chegada de ferramentas de inteligência artificial generativa (IAGen) de acesso público — "
         "ChatGPT (nov. 2022), Gemini e congêneres — expôs a educação superior em Direito no Brasil a um "
         "ponto de inflexão epistemológico, pedagógico e regulatório. A literatura sentinela de Paula (2026) "
         "demonstra que o fenômeno não é periférico: ele tensiona a integridade acadêmica, a avaliação da "
         "aprendizagem e a formação do raciocínio jurídico, exigindo resposta institucional que preserve a "
         "vocação humanista do Direito sem ignorar o paradigma tecnológico emergente.")
add_body(doc,
         "No plano normativo, o Brasil acumula camadas que já conformam o uso educacional de IA, ainda que "
         "de forma fragmentada: a Lei Geral de Proteção de Dados (Lei n. 13.709/2018), o Marco Civil da "
         "Internet, o Projeto de Lei n. 2.338/2023 (Marco Legal da IA, aprovado no Senado em dezembro de "
         "2024 e em tramitação na Câmara dos Deputados em 2026, com 68 obrigações mapeadas pelo ITS Rio), e "
         "a Resolução CNJ n. 615/2025 (DJe n. 54/2025, 14 mar 2025, p. 2-17), que institui governança, "
         "classificação de risco, avaliação de impacto algorítmico obrigatória para alto risco e o catálogo "
         "Sinapses, com capítulo específico para large language models (LLMs) e IAGen (arts. 19-21).")
add_body(doc,
         "Em perspectiva comparada, a União Europeia publicou o AI Act (Regulamento (UE) 2024/1689, JOUE 12 "
         "jul 2024, vigor 01 ago 2024, aplicação faseada até 02 ago 2026/2027), adotando arquitetura "
         "risk-based em quatro níveis (inaceitável, alto, limitado, mínimo) e listando educação e emprego "
         "como domínios de alto risco (Anexo III). Portugal, via Circular da Ordem dos Advogados Portuguesa "
         "(2023), transpõe essa lógica para a deontologia e a formação. Os Estados Unidos mantêm regulação "
         "setorial descentralizada, e a UNESCO (Guidance, 2023, DOI 10.54675/EWZM9535) propõe sete passos "
         "humanocêntricos, idade mínima de 13 anos para uso autônomo e validação ético-pedagógica pelas "
         "instituições (menos de 10% das 450 instituições pesquisadas tinham política para IAGen).")
add_body(doc,
         "O descompasso é o problema: proliferação normativa versus ausência de diretrizes pedagógicas "
         "operacionais, equitativas e auditáveis para a sala de aula jurídica. A revisão sistemática "
         "sentinela de Valencia Jiménez e Beltrán Pacheco (2025), com 19 artigos incluídos (PRISMA, "
         "Scopus/WoS, 136→19), oferece a medida desse descompasso: 57,89% dos estudos concentram-se em "
         "marcos regulatórios, 42,11% em comparação internacional e 42,11% em responsabilidade/rendição de "
         "contas, enquanto acesso à justiça/equidade é sub-representado e justiça social, governança "
         "algorítmica e direitos humanos configuram vácuo crítico (conclusão autoral). Carmo e Alves "
         "(2025), em análise documental de 13 documentos + AI Act/UNESCO/LGPD, confirmam a escassez quando "
         "se cruza ética, tecnologia e ensino jurídico-tributário, propondo o princípio da Transparência "
         "Hermenêutica. Abal e Pilati (2025), com GPTs personalizados (tutoria, mentoria reversa, simulação "
         "de escritório), evidenciam que o ganho de engajamento só se sustenta com letramento crítico — sob "
         "risco de plágio, alucinação e viés.")
add_body(doc,
         "Este manuscrito situa-se no campo da Educação e da educação jurídica e formula a ponte "
         "obrigatória entre IA e Direito via Educação Jurídica e propõe fechar o elo perdido por meio de "
         "síntese secundária rigorosa. A pergunta que orienta o protocolo, no formato PICo comparado, é: "
         "Que evidências, princípios e lacunas a literatura científica (2020-2025) apresenta sobre a "
         "implementação ética, transparente e equitativa de IA Generativa na educação jurídica brasileira "
         "em perspectiva comparada (UE, Portugal, Alemanha, EUA, UNESCO)?")
add_body(doc,
         "A contribuição é tripla: (a) demonstrar com método reprodutível que o GAP não é retórico, mas "
         "mensurável; (b) oferecer matriz comparativa normativa Brasil vs UE/Portugal/Alemanha/EUA/UNESCO "
         "para a discussão pedagógica; (c) propor agenda de pesquisa e princípios operacionais "
         "(Transparência Hermenêutica) contextualizados ao Ensino. Trata-se de revisão de escopo (scoping "
         "review) — desenhada para mapear campo heterogêneo e identificar lacunas — com relato "
         "PRISMA-ScR, com registro OSF previsto antes da execução. Resultados somente serão reportados "
         "após execução e auditoria da busca, triagem e extração.")

# ============ 2 REFERENCIAL TEÓRICO ============
add_heading_custom(doc, "2 REFERENCIAL TEÓRICO")

add_heading_custom(doc, "2.1 IA Generativa, LLMs e educação: o que está em jogo", level=2)
add_body(doc,
         "IAGen e LLMs são sistemas que processam e geram texto, imagem, código e outros artefatos a "
         "partir de grandes bases de dados, por predição estatística (UNESCO, 2023). Na educação, seu "
         "potencial reside em personalização, organização da informação e apoio à argumentação, mas seus "
         "riscos são imediatos: segurança, privacidade, direitos autorais, manipulação, exclusão digital, "
         "dependência tecnológica e perda de controle do processo educacional (UNESCO, 2023; ABAL; PILATI, "
         "2025). A UNESCO alerta que a ausência de regulação nacional na maioria dos países deixa dados "
         "desprotegidos e instituições despreparadas — menos de 10% tinham políticas formais — e que a "
         "validação deve ser dupla: ética e pedagógica, humanocêntrica, com proteção de agência humana, "
         "inclusão, equidade e diversidade linguístico-cultural. O horizonte EdGPT — modelos fundacionais "
         "refinados com conhecimento pedagógico e de interação professor-estudante — ainda depende de "
         "pesquisa robusta sobre coleta ética de dados e sobre não violação de direitos humanos (UNESCO, "
         "2023).")
add_body(doc,
         "No ensino jurídico, a literatura converge para sete tensões: (1) epistemológica (o que conta "
         "como saber jurídico quando a síntese é automatizada), (2) pedagógica (como ensinar raciocínio "
         "jurídico sem terceirizar crítica), (3) avaliativa (como avaliar autoria e integridade), (4) ética "
         "(plágio, alucinação, viés), (5) jurídica (LGPD, direitos autorais no treino), (6) institucional "
         "(quem governa e audita) e (7) equitativa (quem acessa e quem fica para trás).")

add_heading_custom(doc, "2.2 Educação jurídica brasileira: DCNs, currículo e avaliação sob IAGen", level=2)
add_body(doc,
         "As Diretrizes Curriculares Nacionais (DCNs) do curso de Direito exigem compreensão crítica dos "
         "impactos das tecnologias, ética e pensamento crítico. Paula (2026), em perspectiva comparada, e "
         "Silva (2025, UFF, ancorado em Morin), convergem: o fator interveniente decisivo não é a "
         "ferramenta, mas o desenvolvimento de competências e de ética digital. Ghirardi e Feferbaum (FGV, "
         "2023) mapearam práticas discentes preliminares no Brasil, evidenciando uso heterogêneo e demanda "
         "por orientação. A lacuna curricular é, portanto, dupla: ausência de letramento em IA para "
         "docentes/discentes e ausência de desenhos avaliativos que preservem autoria sem proibir uso.")

add_heading_custom(doc, "2.3 Marcos normativos em perspectiva comparada (2020-2025)", level=2)
add_body(doc,
         "Brasil. O PL 2.338/2023 estrutura-se em fundamentos (centralidade da pessoa humana, direitos "
         "humanos, educação para cidadania — art. 2º XIII, proteção de vulneráveis e de "
         "crianças/adolescentes — art. 2º XIV) e princípios (supervisão humana efetiva, não discriminação, "
         "transparência). O art. 14, II, classifica como alto risco sistemas usados como fator determinante "
         "em seleção de estudantes, avaliações determinantes do progresso acadêmico ou monitoramento de "
         "estudantes (ressalvada segurança), com impacto direto na educação jurídica. O capítulo de direitos "
         "do afetado consagra explicação, contestação e revisão humana. Os arts. 29 (avaliação preliminar "
         "para IAGen e modelos de propósito geral) e 25 (avaliação de impacto algorítmico para alto risco) e "
         "o Sistema Nacional de Regulação e Governança (SIA, coord. ANPD, art. 45) completam a arquitetura. "
         "A ITS Rio (23 jun 2025) mapeou 68 obrigações no PL vs 43 no AI Act, com horizontalidade (34/68 "
         "solidárias) e 14 obrigações específicas para o setor público — inovação brasileira.")
add_body(doc,
         "A Resolução CNJ 615/2025 (11 mar 2025) traduz parte dessa arquitetura para o Judiciário: "
         "governança, auditoria e monitoramento proporcionais ao risco (art. 1º §§1º-4º), compatibilidade "
         "com direitos fundamentais (art. 5º), preservação de igualdade e não discriminação (art. 8º), "
         "categorização de risco (art. 9º, Anexo), medidas de governança para alto risco (art. 13), "
         "avaliação de impacto algorítmico (art. 14), Comitê Nacional de IA do Judiciário (art. 15), uso e "
         "contratação de LLMs/SLMs/IAGen (arts. 19-20), indicação em interface de modelos/versão/código "
         "Sinapses (art. 21), conformidade com LGPD/LAI/propriedade intelectual e cadastro no Sinapses com "
         "sumário público de impacto para alto risco e publicação em linguagem simples (arts. 22-25).")
add_quote_long(doc,
         "Art. 14, II, PL 2.338/2023 — Considera-se de alto risco o sistema de IA utilizado como fator "
         "determinante na tomada de decisões de seleção de estudantes em processos de ingresso em "
         "instituições de ensino ou de formação profissional, ou para avaliações determinantes no progresso "
         "acadêmico ou monitoramento de estudantes, ressalvadas as hipóteses de monitoramento "
         "exclusivamente para finalidade de segurança (BRASIL, 2023).")
add_body(doc,
         "União Europeia e Portugal. O AI Act (2024/1689) combina obrigações ex ante (gestão de risco, "
         "governança de dados, documentação, rastreabilidade, supervisão humana) e regime ex post de "
         "responsabilidade civil (Diretiva 2024/2853 + direitos nacionais). Juristech (18 abr 2026) "
         "sintetiza a diferença sistêmica: o PL brasileiro fecha o circuito (conformidade + "
         "responsabilidade objetiva/solidária na mesma lei, arts. 35-36), enquanto a UE separa funções. Em "
         "direitos autorais, o Brasil caminha para regime mais intervencionista (transparência de obras no "
         "treino + opt-out + remuneração) versus exceção de mineração de textos e dados (TDM) com reserva "
         "na UE. Portugal, via OAP Circular 2023, alinha deontologia e formação à transposição do AI Act e "
         "ao RGPD.")
add_body(doc,
         "UNESCO (2023) ancora todo o edifício em abordagem humanocêntrica (Recomendação Ética IA 2021 + "
         "Beijing Consensus 2019), defendendo regulação que garanta agência humana, transparência e "
         "accountability pública, com extensão para currículo, ensino, aprendizagem e pesquisa.")

add_heading_custom(doc, "2.4 Transparência Hermenêutica como princípio operacional", level=2)
add_body(doc,
         "Carmo e Alves (2025) avançam além da transparência algorítmica geral (informacional, art. 6º VI "
         "LGPD; publicidade art. 37 CF; fundamentação art. 93 IX CF) ao propor Transparência Hermenêutica "
         "como categoria dogmática autônoma: o sistema deve declarar qual método interpretativo adotou, em "
         "que base axiológica o fundamentou e quais limites sistemáticos — notadamente a tipicidade fechada "
         "do direito tributário (Queiroz; Baleeiro/Derzi) — circunscreveram sua operação. A coincidência "
         "nominal com transparência informacional é apenas aparente: aqui, exige-se auditabilidade do "
         "raciocínio, não apenas do resultado. O Supremo Tribunal Federal, na ADI 6.649/DF, já determinou "
         "explicabilidade de sistemas eletrônicos, fundamento jurisprudencial para o princípio. Em educação "
         "jurídica, a Transparência Hermenêutica operacionaliza-se em rubricas que exigem do estudante (e do "
         "docente que usa IAGen): declarar método, base axiológica e limites dentro dos quais a ferramenta "
         "operou.")

# ============ 3 MÉTODO ============
add_heading_custom(doc, "3 MÉTODO")

add_heading_custom(doc, "3.1 Desenho e diretriz de relato", level=2)
add_body(doc,
         "Será conduzida revisão de escopo orientada pela metodologia do Joanna Briggs Institute, com "
         "relato segundo PRISMA-ScR. A escolha decorre da heterogeneidade esperada do campo, que reúne "
         "estudos pedagógicos, análises documentais, relatos de experiência e literatura normativa. O "
         "objetivo é mapear conceitos, tipos de evidência e lacunas, e não estimar um efeito agregado.")

add_heading_custom(doc, "3.2 Protocolo e registro", level=2)
add_body(doc,
         "O protocolo deve ser registrado no Open Science Framework antes da execução definitiva da "
         "revisão. Na data desta versão, o arquivo fornecido contém apenas o campo de registro a criar; "
         "portanto, nenhum DOI ou URL de registro é declarado como existente. Strings, critérios de "
         "elegibilidade, formulário de extração e eventuais emendas deverão ser versionados.")

add_heading_custom(doc, "3.3 Pergunta e escopo", level=2)
add_body(doc,
         "Pergunta principal: que evidências, princípios e lacunas a literatura científica de 2020–2025 "
         "apresenta sobre a implementação ética, transparente e equitativa de IAGen na educação jurídica "
         "brasileira em perspectiva comparada? População/contexto: educação jurídica superior e educação "
         "judicial; conceito: IAGen/LLMs em ensino, aprendizagem e avaliação; contexto comparado: Brasil, "
         "União Europeia, Portugal, Alemanha e Estados Unidos, com documentos internacionais pertinentes.")

add_heading_custom(doc, "3.4 Fontes de informação e busca", level=2)
add_body(doc,
         "Planejam-se buscas em Scopus, Web of Science Core Collection, SciELO, Educ@, Portal de Periódicos "
         "CAPES e DOAJ, com Google Acadêmico apenas como fonte complementar. Cada busca deverá registrar "
         "data, string integral, filtros, quantidade recuperada e arquivo exportado. O recorte principal é "
         "2020–2025; documentos de 2026 poderão ser usados apenas como atualização contextual fora do "
         "corpus, salvo alteração formal e prospectiva do protocolo.")

add_heading_custom(doc, "3.5 Elegibilidade", level=2)
add_body(doc,
         "Incluir artigos originais, revisões e análises documentais com foco explícito em educação "
         "jurídica e IA, publicados no período definido, com texto completo e método identificável. Excluir "
         "duplicatas, opiniões sem método, trabalhos cujo foco seja exclusivamente automação judiciária sem "
         "dimensão educacional e documentos fora do período do corpus. Literatura normativa e institucional "
         "deverá ser distinguida da literatura empírica e não poderá ser contabilizada como estudo "
         "pedagógico sem justificativa metodológica.")

add_heading_custom(doc, "3.6 Seleção e extração", level=2)
add_body(doc,
         "A triagem deverá ser realizada por revisores independentes, com resolução documentada de "
         "divergências. O fluxo deverá registrar identificação, deduplicação, triagem por título/resumo, "
         "avaliação de texto completo, motivos de exclusão e conjunto final incluído. A ficha de extração "
         "deverá conter metadados, país/contexto, desenho, amostra ou corpus, método, achados, limitações, "
         "financiamento/conflitos, dimensões analíticas e lições transferíveis ao Brasil.")

add_heading_custom(doc, "3.7 Síntese", level=2)
add_body(doc,
         "A síntese quantitativa será descritiva, com denominadores explícitos. Frequências por dimensão "
         "somente serão calculadas a partir da matriz validada. A síntese qualitativa organizará "
         "convergências, divergências e lacunas, mantendo separados resultados dos estudos, interpretação "
         "dos autores desta revisão e proposições normativas.")

add_heading_custom(doc, "3.8 Integridade, ética e transparência", level=2)
add_body(doc,
         "Por utilizar literatura e documentos públicos, a revisão não envolve recrutamento de "
         "participantes. A necessidade de apreciação ética deverá ser confirmada conforme as regras "
         "institucionais aplicáveis. Financiamento, conflitos, contribuições autorais, disponibilidade de "
         "dados e uso de ferramentas de IA deverão ser declarados de forma específica e verdadeira. Não se "
         "adota percentual universal de similaridade como garantia de originalidade.")

# ============ 4 RESULTADOS ============
add_heading_custom(doc, "4 RESULTADOS")
add_body(doc, "[A CONFIRMAR APÓS EXECUÇÃO REAL E AUDITÁVEL DA REVISÃO].",
         bold=True, size=12, first_line=False, alignment=WD_ALIGN_PARAGRAPH.CENTER)
add_body(doc,
         "Esta versão não apresenta contagens de registros identificados, duplicatas, exclusões, estudos "
         "incluídos ou frequências temáticas como resultados observados. Os arquivos de trabalho fornecidos "
         "descrevem a extração Rayyan como simulada; por isso, os valores n=1.042, n=15, 66,7%, 60,0%, "
         "53,3%, 46,7%, 33,3%, 26,7% e 20,0%, bem como o radar e o gap map derivados desses valores, foram "
         "excluídos do corpo científico desta versão. Eles somente poderão retornar ao manuscrito se forem "
         "reproduzidos a partir de exportações reais das bases, histórico de deduplicação/triagem e matriz "
         "de extração verificável.")

add_heading_custom(doc, "4.1 Estrutura prevista para apresentação dos resultados", level=2)
add_body(doc,
         "Após a validação, esta seção deverá conter: (a) fluxograma PRISMA-ScR com números rastreáveis; "
         "(b) caracterização dos estudos incluídos; (c) frequências das dimensões analíticas com numerador "
         "e denominador; (d) síntese temática; (e) matriz comparativa Brasil–contextos internacionais; (f) "
         "mapa de lacunas. Cada tabela e figura deverá apontar diretamente para a matriz de extração que "
         "lhe dá origem.")

add_heading_custom(doc, "4.2 Figuras 2 e 3", level=2)
add_body(doc,
         "As Figuras 2 (radar) e 3 (gap map) recebidas foram mantidas fora deste manuscrito porque seus "
         "próprios arquivos de origem vinculam os valores a uma extração simulada. Não devem ser submetidas "
         "como evidência observada até validação independente dos dados.")

# ============ 5 DISCUSSÃO ============
add_heading_custom(doc, "5 DISCUSSÃO — ESTRUTURA ANALÍTICA A SER TESTADA")
add_body(doc,
         "A discussão definitiva dependerá dos resultados validados. Nesta fase, quatro proposições "
         "orientam a leitura sem serem tratadas como achados: (1) a literatura pode privilegiar análise "
         "normativa em relação à avaliação pedagógica; (2) equidade e governança podem exigir dados "
         "primários e, por isso, aparecer menos em estudos documentais; (3) experiências internacionais "
         "podem oferecer instrumentos transferíveis, mas sua aplicação ao Brasil exige compatibilidade com "
         "LGPD, regulação educacional e desenho institucional; (4) transparência, supervisão humana e "
         "letramento em IA constituem dimensões a serem examinadas empiricamente, e não presumidas como "
         "consenso.")

add_heading_custom(doc, "5.1 Da norma à prática pedagógica", level=2)
add_body(doc,
         "Fontes normativas e orientações institucionais podem informar critérios de governança, mas não "
         "demonstram, por si, eficácia pedagógica. A análise deverá distinguir obrigações jurídicas, "
         "recomendações éticas e evidências de aprendizagem. Relatos de experiência, quando incluídos, "
         "deverão ser interpretados segundo seu desenho, contexto e limitações.")

add_heading_custom(doc, "5.2 Perspectiva comparada", level=2)
add_body(doc,
         "A comparação internacional deverá evitar transplantes normativos automáticos. Diferenças entre "
         "regimes jurídicos, instituições de ensino, formação profissional e infraestrutura digital precisam "
         "ser explicitadas antes de qualquer recomendação ao contexto brasileiro.")

add_heading_custom(doc, "5.3 Equidade e governança algorítmica", level=2)
add_body(doc,
         "A revisão deverá verificar como os estudos definem e mensuram equidade, acesso, viés, "
         "explicabilidade, contestação e supervisão humana. A ausência de medidas empíricas não poderá ser "
         "convertida em prova de inexistência do fenômeno; deverá ser reportada como lacuna de evidência.")

# ============ 6 CONCLUSÃO PROVISÓRIA ============
add_heading_custom(doc, "6 CONCLUSÃO PROVISÓRIA")
add_body(doc,
         "O manuscrito apresenta uma arquitetura metodológica adequada a uma revisão de escopo sobre IAGen "
         "e educação jurídica, mas ainda não sustenta conclusões empíricas sobre frequência de dimensões ou "
         "magnitude de lacunas. A contribuição potencial reside em integrar literatura pedagógica, jurídica "
         "e de governança com rastreabilidade metodológica. A conclusão final deverá ser reescrita somente "
         "após a execução real da busca, triagem e extração, preservando linguagem proporcional aos desenhos "
         "encontrados.")

# ============ DECLARAÇÕES ============
add_heading_custom(doc, "DECLARAÇÕES")
add_body(doc, "Financiamento: [A CONFIRMAR].", first_line=False)
add_body(doc, "Conflitos de interesses: [A CONFIRMAR PELOS AUTORES].", first_line=False)
add_body(doc, "Contribuições CRediT: [A CONFIRMAR COM BASE NAS CONTRIBUIÇÕES REAIS].", first_line=False)
add_body(doc,
         "Disponibilidade de dados: protocolo OSF [A REGISTRAR]; matriz de extração e exportações das "
         "bases [A VALIDAR E DEPOSITAR].", first_line=False)
add_body(doc,
         "Uso de inteligência artificial: [A PREENCHER DE FORMA ESPECÍFICA, INDICANDO FERRAMENTA, "
         "FINALIDADE, ETAPAS E SUPERVISÃO HUMANA].", first_line=False)
add_body(doc,
         "Autoria/ORCID/afiliação: devem constar apenas na folha de identificação separada, conforme as "
         "instruções atuais do periódico-alvo.", first_line=False)

# ============ REFERÊNCIAS ============
add_heading_custom(doc, "REFERÊNCIAS — NÚCLEO VERIFICADO/PENDENTE DE AUDITORIA INTEGRAL")
refs = [
    "ABAL, Felipe Cittolin; PILATI, Adriana Fasolo. Inteligência artificial no ensino jurídico: "
    "experiências com GPTs personalizados. International Journal of Digital Law, v. 6, e609, 2025. "
    "DOI: 10.47975/ijdl.v.6.1302.",

    "BRASIL. Lei n. 13.709, de 14 de agosto de 2018. Lei Geral de Proteção de Dados Pessoais.",

    "BRASIL. Senado Federal. Projeto de Lei n. 2.338, de 2023. Dispõe sobre o uso da inteligência "
    "artificial. [Conferir situação legislativa na data da submissão].",

    "BRASIL. Conselho Nacional de Justiça. Resolução n. 615, de 11 de março de 2025. Estabelece "
    "diretrizes para o desenvolvimento, utilização e governança de soluções de inteligência artificial no "
    "Poder Judiciário.",

    "PAGE, Matthew J. et al. The PRISMA 2020 statement: an updated guideline for reporting systematic "
    "reviews. BMJ, v. 372, n. 71, 2021. DOI: 10.1136/bmj.n71.",

    "PETERS, Micah D. J. et al. Updated methodological guidance for the conduct of scoping reviews. "
    "JBI Evidence Synthesis, v. 18, n. 10, p. 2119–2126, 2020. DOI: 10.11124/JBIES-20-00167.",

    "TRICCO, Andrea C. et al. PRISMA Extension for Scoping Reviews (PRISMA-ScR): checklist and "
    "explanation. Annals of Internal Medicine, v. 169, p. 467–473, 2018.",

    "UNESCO; MIAO, Fengchun; HOLMES, Wayne. Guidance for generative AI in education and research. "
    "Paris: UNESCO, 2023. DOI: 10.54675/EWZM9535.",

    "UNIÃO EUROPEIA. Regulamento (UE) 2024/1689 do Parlamento Europeu e do Conselho, de 13 de junho de "
    "2024, que estabelece regras harmonizadas em matéria de inteligência artificial.",

    "VALENCIA JIMÉNEZ, Carla Alessandra; BELTRÁN PACHECO, Jorge Alberto. Dilemas éticos de la "
    "Inteligencia Artificial en la práctica legal: Revisión Sistemática. Revista Tribunal, v. 5, n. 12, "
    "p. 748–766, 2025. DOI: 10.59659/revistatribunal.v5i12.235.",

    "PAULA, Gil César Costa de. Inteligência artificial generativa: desafios para o ensino jurídico "
    "superior brasileiro. Revista OWL, v. 4, n. 5, p. 1–22, 2026. DOI: 10.5281/zenodo.19966396. "
    "[Usar apenas como atualização contextual se o corpus permanecer 2020–2025].",
]
for r in refs:
    add_reference(doc, r)

# ============ NOTA DE AUDITORIA EDITORIAL ============
add_heading_custom(doc, "NOTA DE AUDITORIA EDITORIAL")
add_body(doc,
         "O periódico Educação Por Escrito é atualmente classificado como Qualis A4 na área principal, e "
         "não A1. Assim, “Qualis A1” pode ser tratado apenas como meta de rigor ou como requisito para "
         "seleção de outro periódico. As instruções atuais da Educação Por Escrito exigem, entre outros "
         "pontos, manuscrito desidentificado, folha de identificação separada, declaração de uso de IA, "
         "ORCID, formato Word e extensão de 15 a 20 páginas. Antes de qualquer submissão, a revista-alvo e "
         "sua classificação devem ser confirmadas novamente na fonte oficial vigente.",
         size=9, italic=True, space_after=12)

doc.save(OUTPUT)
print(f"DOCX v2.0 (PROTOCOLO) salvo: {OUTPUT}")
print(f"Parágrafos: {len(doc.paragraphs)}, Tabelas: {len(doc.tables)}")